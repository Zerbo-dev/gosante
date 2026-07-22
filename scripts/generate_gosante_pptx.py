#!/usr/bin/env python3
"""Génère la présentation PowerPoint GoSanté pour porteurs non tech."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from copy import deepcopy
from lxml import etree

OUT = Path(__file__).resolve().parents[1] / "docs" / "GoSante_Presentation_Porteurs.pptx"

TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_DARK = RGBColor(0x11, 0x5E, 0x59)
TEAL_SOFT = RGBColor(0xCC, 0xFB, 0xF1)
SLATE = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
AMBER = RGBColor(0xC2, 0x41, 0x0C)


def set_run(run, size=18, bold=False, color=SLATE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=SLATE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=color)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    return shape


def add_round_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    try:
        shape.adjustments[0] = 0.15
    except Exception:
        pass
    return shape


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def footer(slide, text="GoSanté · Confidentiel partenaires · Juillet 2026"):
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.05),
        Inches(9),
        Inches(0.3),
        text,
        size=10,
        color=MUTED,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, TEAL_DARK)
    add_rect(s, 0, Inches(2.4), prs.slide_width, Inches(2.4), TEAL)
    add_textbox(s, Inches(0.8), Inches(2.7), Inches(11.5), Inches(1), "GoSanté", size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(
        s,
        Inches(1),
        Inches(3.7),
        Inches(11.3),
        Inches(0.7),
        "La santé plus proche — plateforme numérique pour le Burkina Faso",
        size=22,
        color=TEAL_SOFT,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        Inches(1),
        Inches(6.4),
        Inches(11.3),
        Inches(0.4),
        "Présentation porteurs de projet · Non technique · Juillet 2026",
        size=14,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )

    # 2. Agenda
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Au programme", size=32, bold=True, color=TEAL_DARK)
    add_bullets(
        s,
        Inches(0.9),
        Inches(1.3),
        Inches(11),
        Inches(5.5),
        [
            "Le problème et la promesse",
            "Pour qui ? Les 5 rôles",
            "Les modules principaux",
            "L’expérience utilisateur",
            "Comment ça marche (sans jargon)",
            "Confiance & données",
            "Où nous en sommes + prochaines étapes",
            "Essayer la démo",
        ],
        size=20,
    )
    footer(s)

    # 3. Problem
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Le problème sur le terrain", size=32, bold=True, color=TEAL_DARK)
    cards = [
        ("Médicaments", "Trouver une pharmacie, surtout de garde, et savoir ce qui est disponible."),
        ("Dossier fragile", "Informations santé souvent papier, incomplètes, difficiles à partager."),
        ("Soutien limité", "Accès restreint à l’écoute psychologique et à l’orientation d’urgence."),
        ("Coordination", "Pharmacie, patient et livreur peinent à synchroniser une livraison fiable."),
    ]
    for i, (title, body) in enumerate(cards):
        x = Inches(0.6) + (i % 2) * Inches(6.2)
        y = Inches(1.3) + (i // 2) * Inches(2.6)
        add_round_rect(s, x, y, Inches(5.9), Inches(2.3), LIGHT)
        add_textbox(s, x + Inches(0.3), y + Inches(0.35), Inches(5.3), Inches(0.5), title, size=20, bold=True, color=TEAL)
        add_textbox(s, x + Inches(0.3), y + Inches(1.0), Inches(5.3), Inches(1.0), body, size=15, color=MUTED)
    footer(s)

    # 4. Solution
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "La solution GoSanté", size=32, bold=True, color=TEAL_DARK)
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.2),
        Inches(12),
        Inches(0.8),
        "Une application web mobile-first qui regroupe les services utiles du quotidien santé.",
        size=20,
        color=SLATE,
    )
    pillars = [
        ("Pharmacie", "Carte, catalogue, commande, livraison"),
        ("Bien-être", "Journal, exercices, chat d’écoute"),
        ("Carnet", "Dossier personnel + export PDF"),
        ("Consultation", "RDV psychologue + visio"),
    ]
    for i, (t, b) in enumerate(pillars):
        x = Inches(0.55) + i * Inches(3.15)
        add_round_rect(s, x, Inches(2.4), Inches(3.0), Inches(3.4), TEAL)
        add_textbox(s, x + Inches(0.2), Inches(2.9), Inches(2.6), Inches(0.8), t, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(s, x + Inches(0.2), Inches(3.9), Inches(2.6), Inches(1.5), b, size=15, color=TEAL_SOFT, align=PP_ALIGN.CENTER)
    footer(s)

    # 5. Roles
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Pour qui ? Cinq rôles", size=32, bold=True, color=TEAL_DARK)
    roles = [
        ("Patient", "Commander, suivre, carnet, psy, bien-être"),
        ("Pharmacien", "Stock et préparation des commandes"),
        ("Livreur", "Missions GPS + code de remise"),
        ("Psychologue", "Créneaux, RDV, visio, notes"),
        ("Admin", "Pilotage, rôles, assignations"),
    ]
    for i, (t, b) in enumerate(roles):
        y = Inches(1.2) + i * Inches(1.05)
        add_round_rect(s, Inches(0.7), y, Inches(11.9), Inches(0.9), LIGHT if i % 2 == 0 else TEAL_SOFT)
        add_textbox(s, Inches(1.0), y + Inches(0.2), Inches(3), Inches(0.5), t, size=18, bold=True, color=TEAL_DARK)
        add_textbox(s, Inches(4.2), y + Inches(0.22), Inches(8), Inches(0.5), b, size=16, color=SLATE)
    footer(s)

    # 6. Pharmacy deep dive
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Module phare — Pharmacie", size=32, bold=True, color=TEAL_DARK)
    add_bullets(
        s,
        Inches(0.8),
        Inches(1.3),
        Inches(7.5),
        Inches(5.2),
        [
            "218 pharmacies sur une carte claire et utilisable",
            "Filtres : ville, de garde, ouvertes maintenant",
            "794 médicaments LNME 2023 + prix de référence 2025",
            "Panier, commande, urgence, point GPS de livraison",
            "Suivi jusqu’à la remise avec code patient",
            "Espace pharmacien (stock) + espace livreur (missions)",
        ],
        size=18,
    )
    add_round_rect(s, Inches(8.7), Inches(1.5), Inches(4.0), Inches(4.5), TEAL_DARK)
    add_textbox(s, Inches(9.0), Inches(2.1), Inches(3.4), Inches(0.6), "En un parcours", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(
        s,
        Inches(9.0),
        Inches(2.9),
        Inches(3.4),
        Inches(2.6),
        "Chercher → Commander → Préparer → Livrer → Confirmer",
        size=16,
        color=TEAL_SOFT,
        align=PP_ALIGN.CENTER,
    )
    footer(s)

    # 7. Mental + consult
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Bien-être & consultation", size=32, bold=True, color=TEAL_DARK)
    add_round_rect(s, Inches(0.6), Inches(1.3), Inches(5.9), Inches(5.0), LIGHT)
    add_textbox(s, Inches(0.9), Inches(1.55), Inches(5.4), Inches(0.5), "Santé mentale", size=22, bold=True, color=TEAL)
    add_bullets(
        s,
        Inches(0.9),
        Inches(2.3),
        Inches(5.3),
        Inches(3.6),
        [
            "Journal d’humeur",
            "Coffre chiffré personnel",
            "Chat d’écoute encadré",
            "Exercices guidés",
            "Numéros d’urgence BF",
        ],
        size=16,
    )
    add_round_rect(s, Inches(6.8), Inches(1.3), Inches(5.9), Inches(5.0), LIGHT)
    add_textbox(s, Inches(7.1), Inches(1.55), Inches(5.4), Inches(0.5), "Psychologue & visio", size=22, bold=True, color=TEAL)
    add_bullets(
        s,
        Inches(7.1),
        Inches(2.3),
        Inches(5.3),
        Inches(3.6),
        [
            "Annuaire et créneaux",
            "RDV avec option anonyme",
            "Visio dans l’application",
            "Notes de séance côté pro",
            "Avis patients",
        ],
        size=16,
    )
    footer(s)

    # 8. UX
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Expérience utilisateur", size=32, bold=True, color=TEAL_DARK)
    ux = [
        ("Mobile d’abord", "Conçu pour smartphone, utilisable aussi sur ordinateur."),
        ("Cartes utiles", "Marqueurs clairs, légende, recentrage, itinéraire visible."),
        ("Parcours guidés", "Peu d’étapes pour commander ou prendre un RDV."),
        ("Installable", "PWA : GoSanté s’ajoute à l’écran d’accueil."),
    ]
    for i, (t, b) in enumerate(ux):
        x = Inches(0.55) + (i % 4) * Inches(3.15)
        add_round_rect(s, x, Inches(1.5), Inches(3.0), Inches(4.4), TEAL if i % 2 == 0 else TEAL_DARK)
        add_textbox(s, x + Inches(0.2), Inches(2.2), Inches(2.6), Inches(1.0), t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(s, x + Inches(0.2), Inches(3.5), Inches(2.6), Inches(1.8), b, size=14, color=TEAL_SOFT, align=PP_ALIGN.CENTER)
    footer(s)

    # 9. Stack simple
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Comment ça marche (sans jargon)", size=30, bold=True, color=TEAL_DARK)
    rows = [
        ("Application web", "Les écrans que voient les utilisateurs"),
        ("Supabase", "Comptes, données, sécurité, temps réel"),
        ("Vercel", "Mise en ligne de la plateforme"),
        ("Cartes libres", "Pharmacies, trajets, livraison"),
        ("IA encadrée", "Soutien conversationnel (pas un médecin)"),
        ("Visio", "Consultation à distance dans le navigateur"),
    ]
    for i, (t, b) in enumerate(rows):
        y = Inches(1.15) + i * Inches(0.9)
        add_round_rect(s, Inches(0.7), y, Inches(11.9), Inches(0.8), LIGHT if i % 2 == 0 else TEAL_SOFT)
        add_textbox(s, Inches(1.0), y + Inches(0.18), Inches(3.5), Inches(0.45), t, size=16, bold=True, color=TEAL_DARK)
        add_textbox(s, Inches(4.8), y + Inches(0.18), Inches(7.5), Inches(0.45), b, size=16, color=SLATE)
    footer(s)

    # 10. Trust
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Confiance & protection des données", size=30, bold=True, color=TEAL_DARK)
    add_bullets(
        s,
        Inches(0.9),
        Inches(1.3),
        Inches(11.5),
        Inches(5.2),
        [
            "Consentement obligatoire à l’inscription",
            "Politique de confidentialité accessible",
            "Chaque rôle ne voit que ce qui le concerne",
            "Traçabilité des actions importantes (audit)",
            "Journal mental chiffrable par l’utilisateur",
            "Cadre légal BF (Loi 001-2021/AN) — conformité hébergement à finaliser",
        ],
        size=20,
    )
    footer(s)

    # 11. Status
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Où en sommes-nous ?", size=32, bold=True, color=TEAL_DARK)
    status = [
        ("Prêt à démontrer", "Pharmacie bout-en-bout, livraison, santé mentale, carnet PDF, visio psy, admin"),
        ("En consolidation", "Paiement réel Mobile Money, scan ordonnance dans le menu, PWA"),
        ("À venir", "CGU, conformité hébergement BF, pilote pharmacies, Phase 3 écosystème"),
    ]
    colors = [TEAL, TEAL_DARK, AMBER]
    for i, ((t, b), c) in enumerate(zip(status, colors)):
        x = Inches(0.55) + i * Inches(4.2)
        add_round_rect(s, x, Inches(1.5), Inches(4.0), Inches(4.5), c)
        add_textbox(s, x + Inches(0.25), Inches(1.9), Inches(3.5), Inches(1.0), t, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(s, x + Inches(0.25), Inches(3.2), Inches(3.5), Inches(2.2), b, size=15, color=WHITE, align=PP_ALIGN.CENTER)
    footer(s)

    # 12. Next steps
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(0.15), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.6), "Prochaines étapes recommandées", size=30, bold=True, color=TEAL_DARK)
    add_bullets(
        s,
        Inches(0.9),
        Inches(1.3),
        Inches(11.5),
        Inches(5.2),
        [
            "Brancher le paiement Mobile Money réel",
            "Réactiver le scan d’ordonnance dans le menu",
            "Valider l’hébergement des données de santé au Burkina Faso",
            "Rédiger CGU + guide utilisateur simple",
            "Lancer un pilote avec pharmacies et livreurs partenaires",
        ],
        size=20,
    )
    footer(s)

    # 13. CTA / Demo
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, TEAL_DARK)
    add_textbox(s, Inches(1), Inches(2.0), Inches(11.3), Inches(0.8), "Essayer la démo", size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(
        s,
        Inches(1),
        Inches(3.0),
        Inches(11.3),
        Inches(0.7),
        "https://gosantetest.vercel.app",
        size=28,
        bold=True,
        color=TEAL_SOFT,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        Inches(1.5),
        Inches(4.2),
        Inches(10.3),
        Inches(1.2),
        "Créer un compte patient, explorer la carte pharmacies,\nla santé mentale, le carnet et la consultation.",
        size=18,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s,
        Inches(1),
        Inches(6.4),
        Inches(11.3),
        Inches(0.4),
        "Merci — GoSanté, pour une santé plus accessible au Burkina Faso",
        size=14,
        color=TEAL_SOFT,
        align=PP_ALIGN.CENTER,
    )

    prs.save(str(OUT))
    print(f"PPTX écrit : {OUT}")
    return OUT


if __name__ == "__main__":
    build()
