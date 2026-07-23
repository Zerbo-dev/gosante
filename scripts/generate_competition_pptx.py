#!/usr/bin/env python3
"""
Pitch PowerPoint GoSanté — concours (5 min)
Thème : Concevoir et adapter des technologies innovantes
        pour répondre aux défis de santé au Burkina Faso

DA : « Clinique sahélien » — émeraude profond, fond clair froid,
     accent ocre discret, peu de texte, grandes affirmations.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "GoSante_Pitch_Concours.pptx"
ARTIFACT = Path("/opt/cursor/artifacts/GoSante_Pitch_Concours.pptx")

# --- Palette (évite violet / crème-terracotta / mode sombre) ---
BRAND = RGBColor(0x06, 0x5F, 0x46)       # émeraude profond
BRAND_MID = RGBColor(0x05, 0x96, 0x69)   # émeraude vif
BRAND_SOFT = RGBColor(0xEC, 0xFD, 0xF5)  # émeraude très pâle
INK = RGBColor(0x0F, 0x17, 0x2A)         # ardoise
BODY = RGBColor(0x33, 0x41, 0x55)        # ardoise moyen
MUTE = RGBColor(0x64, 0x74, 0x8B)        # ardoise claire
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xF8, 0xFA, 0xFC)         # fond froid
OCHRE = RGBColor(0xB4, 0x53, 0x09)       # accent Burkina (discret)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color)
    return s


def round_rect(slide, l, t, w, h, color, adj=0.12):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color)
    try:
        s.adjustments[0] = adj
    except Exception:
        pass
    return s


def txt(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf.paragraphs[0].space_before = Pt(0)
        tf.paragraphs[0].space_after = Pt(0)
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def multi(slide, l, t, w, h, lines, size=16, color=BODY, bold=False, spacing=10):
    """lines = list of strings"""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)
    return box


def bullets(slide, l, t, w, h, items, size=17, color=BODY):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        run = p.add_run()
        run.text = f"→  {item}"
        set_run(run, size=size, color=color)
    return box


def footer(slide, page, total=7):
    rect(slide, Inches(0), SLIDE_H - Inches(0.42), SLIDE_W, Inches(0.42), BRAND)
    txt(
        slide,
        Inches(0.5),
        SLIDE_H - Inches(0.36),
        Inches(8),
        Inches(0.3),
        "GoSanté  ·  Technologies innovantes pour la santé au Burkina Faso",
        size=11,
        color=WHITE,
    )
    txt(
        slide,
        SLIDE_W - Inches(1.4),
        SLIDE_H - Inches(0.36),
        Inches(1),
        Inches(0.3),
        f"{page}/{total}",
        size=11,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )


def side_bar(slide):
    rect(slide, Inches(0), Inches(0), Inches(0.18), SLIDE_H, BRAND)


def blank(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 7

    # ========== 1. TITLE ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WASH)
    rect(s, Inches(0), Inches(0), Inches(0.35), SLIDE_H, BRAND)
    rect(s, Inches(0), SLIDE_H - Inches(1.15), SLIDE_W, Inches(1.15), BRAND)

    txt(s, Inches(0.8), Inches(1.1), Inches(11), Inches(0.4),
        "CONCOURS  ·  SANTÉ & INNOVATION", size=14, bold=True, color=OCHRE)
    txt(s, Inches(0.8), Inches(1.7), Inches(11.5), Inches(1.6),
        "Concevoir et adapter des technologies\ninnovantes pour la santé au Burkina Faso",
        size=36, bold=True, color=INK)
    txt(s, Inches(0.8), Inches(3.7), Inches(10), Inches(0.6),
        "GoSanté — la santé numérique pensée pour le terrain burkinabè",
        size=20, color=BODY)
    txt(s, Inches(0.8), Inches(4.6), Inches(8), Inches(0.4),
        "Pitch 5 minutes", size=14, color=MUTE)
    txt(s, Inches(0.8), SLIDE_H - Inches(0.85), Inches(10), Inches(0.35),
        "Une réponse concrète aux défis d’accès aux soins, aux médicaments et au suivi patient",
        size=14, color=WHITE)

    # ========== 2. LE DÉFI ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WASH)
    side_bar(s)
    footer(s, 2, total)

    txt(s, Inches(0.7), Inches(0.4), Inches(11), Inches(0.35),
        "01  —  LE DÉFI", size=13, bold=True, color=OCHRE)
    txt(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.9),
        "Au Burkina Faso, le parcours de soins\nest encore trop fragmenté.",
        size=32, bold=True, color=INK)

    cards = [
        ("Pharmacie", "Trouver une pharmacie de garde\net un médicament disponible\nreste un parcours du combattant."),
        ("Suivi", "Le carnet papier se perd.\nL’historique médical ne suit\npas le patient."),
        ("Accès", "Consultation psy, ordonnance,\nlivraison : tout est dispersé,\npeu digitalisé, peu relié."),
    ]
    x0 = Inches(0.7)
    for i, (title, body) in enumerate(cards):
        x = x0 + Inches(i * 4.05)
        round_rect(s, x, Inches(2.6), Inches(3.8), Inches(3.5), WHITE, 0.08)
        rect(s, x, Inches(2.6), Inches(3.8), Inches(0.12), BRAND_MID)
        txt(s, x + Inches(0.3), Inches(2.95), Inches(3.2), Inches(0.4),
            title, size=18, bold=True, color=BRAND)
        multi(s, x + Inches(0.3), Inches(3.5), Inches(3.2), Inches(2.2),
              body.split("\n"), size=15, color=BODY, spacing=6)

    # ========== 3. NOTRE PARTI PRIS ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WASH)
    side_bar(s)
    footer(s, 3, total)

    txt(s, Inches(0.7), Inches(0.4), Inches(11), Inches(0.35),
        "02  —  NOTRE PARTI PRIS", size=13, bold=True, color=OCHRE)
    txt(s, Inches(0.7), Inches(0.9), Inches(12), Inches(1.1),
        "Pas une appli santé générique.\nUne technologie adaptée au contexte.",
        size=30, bold=True, color=INK)

    points = [
        ("Adapter, pas copier", "On part des usages réels : garde des pharmacies, stock local, Mobile Money, faible bande passante."),
        ("Une seule porte d’entrée", "Patient, pharmacien, livreur, psychologue : un même fil pour le parcours de soins."),
        ("Preuve avant partenariats", "Démo opérationnelle avec stocks pilotes, garde 2026 et parcours commande testable."),
    ]
    y = Inches(2.5)
    for title, desc in points:
        round_rect(s, Inches(0.7), y, Inches(11.9), Inches(1.15), WHITE, 0.06)
        rect(s, Inches(0.7), y, Inches(0.14), Inches(1.15), BRAND_MID)
        txt(s, Inches(1.15), y + Inches(0.22), Inches(11), Inches(0.35),
            title, size=18, bold=True, color=INK)
        txt(s, Inches(1.15), y + Inches(0.58), Inches(11), Inches(0.4),
            desc, size=15, color=BODY)
        y += Inches(1.3)

    # ========== 4. LA SOLUTION ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WASH)
    side_bar(s)
    footer(s, 4, total)

    txt(s, Inches(0.7), Inches(0.4), Inches(11), Inches(0.35),
        "03  —  LA SOLUTION", size=13, bold=True, color=OCHRE)
    txt(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.7),
        "GoSanté : la plateforme santé du quotidien.",
        size=30, bold=True, color=INK)

    modules = [
        ("Pharmacie", "Carte, garde, stock,\ncommande & livraison"),
        ("Carnet", "Dossier patient PDF,\nhistorique, urgence"),
        ("Santé mentale", "RDV, visio sécurisée,\njournal protégé"),
        ("Logistique", "Livreurs en ligne,\ncode de livraison,\nsuivi"),
    ]
    for i, (title, body) in enumerate(modules):
        x = Inches(0.7) + Inches(i * 3.1)
        round_rect(s, x, Inches(2.2), Inches(2.9), Inches(3.6), BRAND if i == 0 else WHITE, 0.08)
        title_c = WHITE if i == 0 else BRAND
        body_c = RGBColor(0xD1, 0xFA, 0xE5) if i == 0 else BODY
        txt(s, x + Inches(0.25), Inches(2.55), Inches(2.4), Inches(0.45),
            title, size=18, bold=True, color=title_c)
        multi(s, x + Inches(0.25), Inches(3.25), Inches(2.4), Inches(2.2),
              body.split("\n"), size=15, color=body_c, spacing=8)

    # ========== 5. ADAPTER LA TECH ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WASH)
    side_bar(s)
    footer(s, 5, total)

    txt(s, Inches(0.7), Inches(0.4), Inches(11), Inches(0.35),
        "04  —  ADAPTER L’INNOVATION", size=13, bold=True, color=OCHRE)
    txt(s, Inches(0.7), Inches(0.9), Inches(12), Inches(0.8),
        "Innover, c’est rendre utile ici.",
        size=30, bold=True, color=INK)

    left = [
        "Garde officielle 2026 intégrée (groupes I–IV)",
        "Recherche médicaments sur le stock réel, pas une liste théorique",
        "Import Excel pour le pharmacien (simple, terrain)",
        "Carte + pin GPS = adresse de livraison",
    ]
    right = [
        "Paiement Mobile Money (sandbox) prêt pour le BF",
        "Visio adaptée Safari / iPhone",
        "Carnet médical exportable en PDF professionnel",
        "Rôles patient / pharmacien / livreur / psy",
    ]

    round_rect(s, Inches(0.7), Inches(2.1), Inches(5.8), Inches(4.2), WHITE, 0.06)
    txt(s, Inches(1.0), Inches(2.35), Inches(5.2), Inches(0.4),
        "Ancré dans le réel", size=16, bold=True, color=BRAND)
    bullets(s, Inches(1.0), Inches(2.9), Inches(5.2), Inches(3.1), left, size=15)

    round_rect(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.2), WHITE, 0.06)
    txt(s, Inches(7.1), Inches(2.35), Inches(5.2), Inches(0.4),
        "Pensé pour l’usage", size=16, bold=True, color=BRAND)
    bullets(s, Inches(7.1), Inches(2.9), Inches(5.2), Inches(3.1), right, size=15)

    # ========== 6. PREUVE / TRAJECTOIRE ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, WASH)
    side_bar(s)
    footer(s, 6, total)

    txt(s, Inches(0.7), Inches(0.4), Inches(11), Inches(0.35),
        "05  —  OÙ NOUS EN SOMMES", size=13, bold=True, color=OCHRE)
    txt(s, Inches(0.7), Inches(0.9), Inches(12), Inches(0.7),
        "Une démo vivante, pas une maquette figée.",
        size=28, bold=True, color=INK)

    stats = [
        ("3", "pharmacies pilotes\navec stock démo"),
        ("4", "groupes de garde\n2026 opérationnels"),
        ("1", "parcours complet\ncommande → livraison"),
        ("5", "min pour\nvous le montrer"),
    ]
    for i, (n, label) in enumerate(stats):
        x = Inches(0.7) + Inches(i * 3.1)
        round_rect(s, x, Inches(2.2), Inches(2.9), Inches(3.5), WHITE, 0.08)
        txt(s, x + Inches(0.25), Inches(2.6), Inches(2.4), Inches(0.9),
            n, size=48, bold=True, color=BRAND_MID, align=PP_ALIGN.CENTER)
        multi(s, x + Inches(0.25), Inches(3.8), Inches(2.4), Inches(1.5),
              label.split("\n"), size=15, color=BODY, spacing=4)

    # ========== 7. CLOSING ==========
    s = blank(prs)
    rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, BRAND)
    rect(s, Inches(0), Inches(0), Inches(0.35), SLIDE_H, BRAND_MID)

    txt(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.4),
        "La meilleure tech santé\nn’est pas la plus complexe.",
        size=34, bold=True, color=WHITE)
    txt(s, Inches(0.9), Inches(3.4), Inches(11), Inches(0.8),
        "C’est celle qui s’adapte au Burkina — et qui soigne vraiment le parcours.",
        size=20, color=RGBColor(0xA7, 0xF3, 0xD0))

    round_rect(s, Inches(0.9), Inches(4.6), Inches(5.5), Inches(1.2), BRAND_MID, 0.1)
    txt(s, Inches(1.15), Inches(4.85), Inches(5), Inches(0.35),
        "GoSanté", size=22, bold=True, color=WHITE)
    txt(s, Inches(1.15), Inches(5.3), Inches(5), Inches(0.3),
        "Santé numérique · Ouagadougou", size=14, color=RGBColor(0xD1, 0xFA, 0xE5))

    txt(s, Inches(7.0), Inches(4.9), Inches(5.5), Inches(0.8),
        "Merci.\nDes questions ?",
        size=22, bold=True, color=WHITE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(ARTIFACT))
    print(f"wrote {OUT}")
    print(f"wrote {ARTIFACT}")


if __name__ == "__main__":
    build()
